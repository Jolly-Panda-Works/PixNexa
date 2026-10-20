"""JSON / YAML / TOML / XML / CSV / XLSX."""
import json

import pytest

from pixnexa_engines import ConversionError, Options, registry
from pixnexa_engines.data import as_records, load

SAMPLE = {"app": {"name": "PixNexa", "version": 2, "tags": ["fast", "offline"], "nested": {"ok": True}}}


def write(tmp_path, name, text):
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return p


@pytest.mark.parametrize("target", ["yaml", "toml", "xml", "json"])
def test_json_roundtrip_through_every_format(tmp_path, target):
    src = write(tmp_path, "c.json", json.dumps(SAMPLE))
    out = registry.convert_file(src, target, Options()).outputs[0]
    assert out.suffix == f".{target}"
    back = load(out)
    if target == "xml":                                    # XML has no types: everything comes back as text
        assert back["app"]["name"] == "PixNexa" and back["app"]["tags"] == ["fast", "offline"]
    else:
        assert back == SAMPLE


def test_yaml_and_toml_inputs(tmp_path):
    y = write(tmp_path, "a.yml", "name: Ana\nlist: [1, 2]\n")
    assert json.loads(registry.convert_file(y, "json", Options()).outputs[0].read_text()) == {"name": "Ana",
                                                                                             "list": [1, 2]}
    t = write(tmp_path, "b.toml", 'title = "x"\n[owner]\nname = "Bo"\n')
    assert load(registry.convert_file(t, "yaml", Options()).outputs[0]) == {"title": "x", "owner": {"name": "Bo"}}


def test_xml_attributes_and_text(tmp_path):
    x = write(tmp_path, "d.xml", '<root><item id="1">A</item><item id="2">B</item></root>')
    data = load(x)
    assert data == {"root": {"item": [{"@id": "1", "#text": "A"}, {"@id": "2", "#text": "B"}]}}
    out = registry.convert_file(x, "xml", Options()).outputs[0]
    assert load(out) == data


def test_csv_json_xlsx_tables(tmp_path):
    src = write(tmp_path, "t.csv", "name;age\nAna;31\nBo;25\n")      # semicolons are detected
    j = registry.convert_file(src, "json", Options()).outputs[0]
    assert json.loads(j.read_text()) == [{"name": "Ana", "age": "31"}, {"name": "Bo", "age": "25"}]
    x = registry.convert_file(j, "xlsx", Options()).outputs[0]
    import openpyxl
    ws = openpyxl.load_workbook(x).active
    assert [c.value for c in ws[1]] == ["name", "age"] and ws["A3"].value == "Bo"
    c = registry.convert_file(j, "csv", Options()).outputs[0]
    assert c.read_text(encoding="utf-8-sig").splitlines() == ["name,age", "Ana,31", "Bo,25"]


def test_nested_records_are_flattened():
    rows = as_records({"users": [{"name": "A", "address": {"city": "X"}, "tags": ["a", "b"]}]})
    assert rows == [{"name": "A", "address.city": "X", "tags": "a; b"}]


def test_formula_text_stays_text_in_xlsx(tmp_path):
    src = write(tmp_path, "f.json", json.dumps([{"cell": "=1+1"}]))
    import openpyxl
    ws = openpyxl.load_workbook(registry.convert_file(src, "xlsx", Options()).outputs[0]).active
    assert ws["A2"].value == "=1+1" and ws["A2"].data_type == "s"


def test_not_a_table(tmp_path):
    src = write(tmp_path, "n.json", json.dumps([[1, [2]], {"a": 1}]))
    with pytest.raises(ConversionError, match="not a table"):
        registry.convert_file(src, "csv", Options())


def test_broken_json(tmp_path):
    src = write(tmp_path, "bad.json", "{nope")
    with pytest.raises(ConversionError, match="Could not read"):
        registry.convert_file(src, "yaml", Options())
    assert sorted(p.name for p in tmp_path.iterdir()) == ["bad.json"]


def test_toml_with_nulls_and_list_root(tmp_path):
    src = write(tmp_path, "l.json", json.dumps([{"a": 1, "b": None}]))
    out = registry.convert_file(src, "toml", Options()).outputs[0]
    assert load(out) == {"items": [{"a": 1}]}
