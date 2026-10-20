"""Structured data: JSON, JSON Lines, YAML, TOML, XML <-> each other; CSV / TSV / XLSX for tables."""
from __future__ import annotations

import csv
import datetime as _dt
import io
import json
import re
import sys
from pathlib import Path
from xml.etree import ElementTree as ET

from .base import ConversionError, Options, ProgressCallback, atomic_output, report, target_path

if sys.version_info >= (3, 11):
    import tomllib
else:  # pragma: no cover - Python 3.10
    import tomli as tomllib

INPUT_EXT = (".json", ".jsonl", ".ndjson", ".yaml", ".yml", ".toml", ".xml", ".csv", ".tsv")
TARGETS = ("json", "yaml", "toml", "xml", "csv", "xlsx")
_ATTR, _TEXT = "@", "#text"


# ───────────────────────── reading ─────────────────────────

def _text(path: Path) -> str:
    data = path.read_bytes()
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", "replace")


def _xml_to_obj(elem: ET.Element):
    node: dict = {f"{_ATTR}{k}": v for k, v in elem.attrib.items()}
    for child in elem:
        tag = _local(child.tag)
        value = _xml_to_obj(child)
        if tag in node:
            if not isinstance(node[tag], list):
                node[tag] = [node[tag]]
            node[tag].append(value)
        else:
            node[tag] = value
    text = (elem.text or "").strip()
    if text:
        if not node:
            return text
        node[_TEXT] = text
    return node if node else None


def _local(tag: str) -> str:
    return tag.split("}", 1)[1] if tag.startswith("{") else tag


def load(path) -> object:
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext == ".json":
            return json.loads(_text(path))
        if ext in (".jsonl", ".ndjson"):
            return [json.loads(line) for line in _text(path).splitlines() if line.strip()]
        if ext in (".yaml", ".yml"):
            import yaml
            docs = [d for d in yaml.safe_load_all(_text(path))]
            return docs[0] if len(docs) == 1 else docs
        if ext == ".toml":
            return tomllib.loads(_text(path))
        if ext == ".xml":
            root = ET.fromstring(_text(path).encode("utf-8"))
            return {_local(root.tag): _xml_to_obj(root)}
        if ext in (".csv", ".tsv"):
            return _read_csv(path, "\t" if ext == ".tsv" else None)
        if ext in (".xlsx", ".xlsm"):
            from .office import xlsx_rows
            sheets = xlsx_rows(path)
            if not sheets:
                return []
            return _rows_to_records(next(iter(sheets.values())))
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"Could not read this {ext.lstrip('.').upper()} file: {_short(exc)}") from None
    raise ConversionError(f"Unsupported data file: {ext}")


def _short(exc: Exception) -> str:
    return (str(exc).splitlines() or [exc.__class__.__name__])[0][:160]


def _read_csv(path: Path, delimiter: str | None):
    text = _text(path)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return _rows_to_records(rows)


def _rows_to_records(rows):
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return []
    header = [str(h) if h not in (None, "") else f"column{i + 1}" for i, h in enumerate(rows[0])]
    return [{header[i]: _plain(v) for i, v in enumerate(r) if i < len(header)} for r in rows[1:]]


def _plain(v):
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        return v.isoformat()
    return "" if v is None else v


# ───────────────────────── tables ─────────────────────────

def as_records(data) -> list[dict]:
    """Bring data into rows of {column: value}. Raises ConversionError when it is not table-like."""
    if isinstance(data, dict):
        lists = [v for v in data.values() if isinstance(v, list)]
        if len(data) == 1 and isinstance(next(iter(data.values())), (list, dict)):
            return as_records(next(iter(data.values())))
        if len(lists) == 1 and all(isinstance(x, dict) for x in lists[0]):
            return as_records(lists[0])
        return [_flatten(data)]
    if isinstance(data, list):
        if all(isinstance(x, dict) for x in data):
            return [_flatten(x) for x in data]
        if all(isinstance(x, list) for x in data) and data:
            return _rows_to_records(data)
        if all(not isinstance(x, (dict, list)) for x in data):
            return [{"value": x} for x in data]
    raise ConversionError("This data is not a table, so it cannot become CSV or XLSX.")


def _flatten(d: dict, prefix: str = "") -> dict:
    out = {}
    for k, v in d.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            out.update(_flatten(v, key + "."))
        elif isinstance(v, list):
            out[key] = json.dumps(v, ensure_ascii=False, default=str) if any(
                isinstance(x, (dict, list)) for x in v) else "; ".join("" if x is None else str(x) for x in v)
        else:
            out[key] = _plain(v)
    return out


def _columns(records: list[dict]) -> list[str]:
    cols: list[str] = []
    for r in records:
        for k in r:
            if k not in cols:
                cols.append(k)
    return cols


# ───────────────────────── writing ─────────────────────────

def _to_xml(data) -> bytes:
    if isinstance(data, dict) and len(data) == 1 and not next(iter(data)).startswith(_ATTR):
        tag, value = next(iter(data.items()))
    else:
        tag, value = "root", data
    root = ET.Element(_tag(tag))
    _fill(root, value)
    ET.indent(root)
    return b'<?xml version="1.0" encoding="UTF-8"?>\n' + ET.tostring(root, encoding="utf-8") + b"\n"


def _tag(name) -> str:
    name = re.sub(r"[^\w.\-]", "_", str(name)) or "item"
    return name if re.match(r"[A-Za-z_]", name) else f"_{name}"


def _fill(elem: ET.Element, value):
    if isinstance(value, dict):
        for k, v in value.items():
            if k.startswith(_ATTR):
                elem.set(_tag(k[1:]), _str(v))
            elif k == _TEXT:
                elem.text = _str(v)
            elif isinstance(v, list):
                for item in v:
                    _fill(ET.SubElement(elem, _tag(k)), item)
            else:
                _fill(ET.SubElement(elem, _tag(k)), v)
    elif isinstance(value, list):
        for item in value:
            _fill(ET.SubElement(elem, "item"), item)
    elif value is not None:
        elem.text = _str(value)


def _str(v) -> str:
    if isinstance(v, bool):
        return "true" if v else "false"
    return v.isoformat() if hasattr(v, "isoformat") else str(v)


def _drop_none(v):
    """TOML has no null."""
    if isinstance(v, dict):
        return {str(k): _drop_none(x) for k, x in v.items() if x is not None}
    if isinstance(v, list):
        return [_drop_none(x) for x in v if x is not None]
    return v


def dumps(data, target: str) -> bytes:
    if target == "json":
        return (json.dumps(data, indent=2, ensure_ascii=False, default=str) + "\n").encode("utf-8")
    if target == "yaml":
        import yaml
        return yaml.safe_dump(_jsonable(data), allow_unicode=True, sort_keys=False).encode("utf-8")
    if target == "toml":
        import tomli_w
        value = _drop_none(_jsonable(data, keep_dates=True))
        if not isinstance(value, dict):
            value = {"items": value}
        try:
            return tomli_w.dumps(value).encode("utf-8")
        except (TypeError, ValueError) as exc:
            raise ConversionError(f"This data cannot be written as TOML: {_short(exc)}") from None
    if target == "xml":
        return _to_xml(_jsonable(data))
    if target == "csv":
        records = as_records(data)
        cols = _columns(records)
        buf = io.StringIO()
        writer = csv.DictWriter(buf, fieldnames=cols, lineterminator="\r\n", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)
        return buf.getvalue().encode("utf-8-sig")         # with a BOM, so Excel shows accents correctly
    if target == "xlsx":
        import openpyxl
        records = as_records(data)
        cols = _columns(records)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(cols)
        for r in records:
            ws.append([_xlsx_value(r.get(c)) for c in cols])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"          # text that looks like a formula stays text
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    raise ConversionError(f"Unsupported target format: {target}")


def _xlsx_value(v):
    if isinstance(v, (int, float, bool)) or v is None:
        return v
    return str(v)


def _jsonable(v, keep_dates: bool = False):
    if isinstance(v, dict):
        return {str(k): _jsonable(x, keep_dates) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [_jsonable(x, keep_dates) for x in v]
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        return v if keep_dates else v.isoformat()
    return v


def convert(src, target: str, options: Options, progress: ProgressCallback | None = None) -> list[Path]:
    src = Path(src)
    target = {"yml": "yaml"}.get(target.lower(), target.lower())
    data = load(src)
    report(progress, 0.5)
    payload = dumps(data, target)
    dst = target_path(src, target, options)
    with atomic_output(dst) as temp:
        temp.write_bytes(payload)
    report(progress, 1)
    return [dst]
